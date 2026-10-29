"""Exercise task export downloads through their production HTTP handler."""

import io
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch
import zipfile

from openpyxl import load_workbook
import yaml

from cms.db import (Attachment, Dataset, File, Generator, Manager,
                    ModelSolutionMeta, Statement, Submission, SubtaskValidator,
                    Task, Testcase as DBTestcase)
from cms.server.admin.handlers import HANDLERS
from cms.server.admin.handlers.excel import ExcelExportMixin
from cms.server.admin.handlers.export_handlers import ExportTasksHandler
import tornado.testing
from tornado.web import Application, RequestHandler


class ExportHarness(ExcelExportMixin, RequestHandler):
    get = ExportTasksHandler.get

    def initialize(self, fixture):
        self.fixture = fixture
        self.sql_session = Mock()
        self.sql_session.query.return_value.order_by.return_value.all.return_value = fixture.tasks
        self.service = fixture.service
        self.url = lambda *parts: '/' + '/'.join(parts)

    def get_current_user(self):
        return self.fixture.admin


class TestTasksExport(tornado.testing.AsyncHTTPTestCase):
    # Tornado 4.5 expects runTest when pytest inspects unittest classes.
    def runTest(self):
        pass

    def get_app(self):
        self.tasks = []
        self.admin = SimpleNamespace(permission_all=False)
        self.service = SimpleNamespace(file_cacher=Mock(), add_notification=Mock())
        self.service.file_cacher.get_file_to_path.side_effect = (
            lambda digest, path: Path(path).write_bytes(digest.encode()))
        self.service.file_cacher.get_file.side_effect = (
            lambda digest: io.BytesIO(digest.encode()))
        pattern = next(pattern for pattern, *rest in HANDLERS
                       if rest[0] is ExportTasksHandler)
        return Application([(pattern, ExportHarness, {'fixture': self})],
                           login_url='/login')

    def task(self, task_id=7, name='example', dataset=True):
        task = Task(name=name, title='כותרת')
        task.id = task_id
        Statement(task=task, language='he', digest='hebrew',
                  source_digest='source', source_extension='.docx')
        Statement(task=task, language='en', digest='english')
        if dataset:
            task.active_dataset = Dataset(
                task=task, description='Default', task_type='Batch',
                task_type_parameters=['alone', ['', ''], 'diff'],
                score_type='Sum', score_type_parameters=100)
        self.tasks.append(task)
        return task

    def archive(self, kind):
        response = self.fetch('/tasks/export/' + kind)
        self.assertEqual(response.code, 200)
        self.assertEqual(response.headers['Content-Type'], 'application/zip')
        self.assertIn('tasks_' + kind + '.zip',
                      response.headers['Content-Disposition'])
        return zipfile.ZipFile(io.BytesIO(response.body))

    def test_metadata_uses_admin_id_and_literal_text(self):
        task = self.task(name='=1+1')
        response = self.fetch('/tasks/export/metadata')
        self.assertEqual(response.code, 200)
        self.assertIn('tasks.xlsx', response.headers['Content-Disposition'])
        sheet = load_workbook(io.BytesIO(response.body)).active
        self.assertEqual(list(sheet.values), [
            ('Task name', 'Title', 'ID'), (task.name, task.title, task.id)])
        self.assertEqual(sheet['A2'].data_type, 's')
        self.service.file_cacher.get_file_to_path.assert_not_called()

    def test_statements_all_languages_sources_and_duplicate_names(self):
        self.task(dataset=False)
        self.task(task_id=8, dataset=False)
        with self.archive('statements') as archive:
            self.assertEqual(set(archive.namelist()), {
                f'tasks/{task_id}_example/{filename}'
                for task_id in (7, 8)
                for filename in ('hebrew.pdf', 'hebrew.docx', 'english.pdf')})
            self.assertEqual(archive.read('tasks/7_example/hebrew.docx'), b'source')

    def test_backup_reuses_task_package_export(self):
        self.task()
        self.task(task_id=8)
        with self.archive('backup') as archive:
            self.assertNotIn('tasks/missing_files.json', archive.namelist())
            self.assertNotIn('tasks/INCOMPLETE_BACKUP.txt', archive.namelist())
            for task_id in (7, 8):
                base = f'tasks/{task_id}_example/'
                config = yaml.safe_load(archive.read(base + 'task.yaml'))
                self.assertEqual(config['name'], 'example')
                self.assertEqual(config['title'], 'כותרת')
                self.assertEqual(archive.read(base + 'statements/english.pdf'), b'english')
                with zipfile.ZipFile(io.BytesIO(archive.read(base + 'tests.zip'))) as tests:
                    self.assertEqual(tests.namelist(), [])

    def test_backup_fails_if_any_task_has_no_dataset(self):
        self.task(dataset=False)
        response = self.fetch('/tasks/export/backup', follow_redirects=False)
        self.assertEqual(response.code, 302)
        self.assertEqual(response.headers['Location'], '/tasks')
        self.assertIn('example', self.service.add_notification.call_args.args[2])

    def test_empty_system_exports(self):
        response = self.fetch('/tasks/export/metadata')
        sheet = load_workbook(io.BytesIO(response.body)).active
        self.assertEqual(sheet.max_row, 1)
        for kind in ('statements', 'backup'):
            with self.archive(kind) as archive:
                self.assertEqual(archive.namelist(), [])

    def test_download_requires_login(self):
        self.admin = None
        for kind in ('metadata', 'statements', 'backup'):
            response = self.fetch('/tasks/export/' + kind, follow_redirects=False)
            self.assertEqual(response.code, 302)
            self.assertTrue(response.headers['Location'].startswith('/login'))

    def test_failed_file_read_cleans_temporary_directory(self):
        self.task()
        self.service.file_cacher.get_file_to_path.side_effect = OSError('missing file')
        from cms.server.admin.handlers import export_handlers
        original = export_handlers.tempfile.TemporaryDirectory
        directories = []

        def temporary_directory(**kwargs):
            directory = original(**kwargs)
            directories.append(directory.name)
            return directory

        with patch.object(export_handlers.tempfile, 'TemporaryDirectory', temporary_directory):
            response = self.fetch('/tasks/export/statements', follow_redirects=False)
        self.assertEqual(response.code, 302)
        self.assertTrue(directories)
        self.assertTrue(all(not Path(path).exists() for path in directories))

    def missing_file(self, missing_digest):
        missing = {missing_digest} if isinstance(missing_digest, str) else set(missing_digest)

        def get_file(digest):
            if digest in missing:
                raise KeyError("File not found.")
            return io.BytesIO(digest.encode())

        self.service.file_cacher.get_file.side_effect = get_file

        def retrieve(digest, path):
            if digest in missing:
                # Ensure an incomplete destination cannot appear as a PDF.
                Path(path).write_bytes(b'')
                raise KeyError("File not found.")
            Path(path).write_bytes(digest.encode())
        self.service.file_cacher.get_file_to_path.side_effect = retrieve

    def test_missing_pdf_exports_available_files_and_report(self):
        self.task(dataset=False)
        other = self.task(task_id=8, name='other', dataset=False)
        other.statements['he'].digest = 'other_hebrew'
        self.missing_file('hebrew')
        with self.archive('statements') as archive:
            self.assertNotIn('tasks/7_example/hebrew.pdf', archive.namelist())
            self.assertEqual(archive.read('tasks/7_example/hebrew.docx'), b'source')
            self.assertEqual(archive.read('tasks/7_example/english.pdf'), b'english')
            self.assertEqual(archive.read('tasks/8_other/hebrew.pdf'), b'other_hebrew')
            report = json.loads(archive.read('tasks/missing_statements.json'))
            self.assertEqual(report, [{
                'task_id': 7, 'task_name': 'example', 'language': 'he',
                'kind': 'statement PDF', 'digest': 'hebrew',
            }])
        self.assertIn('missing_statements.json',
                      self.service.add_notification.call_args.args[2])

    def test_missing_source_keeps_pdf_and_reports_source(self):
        self.task()
        self.missing_file('source')
        with self.archive('statements') as archive:
            self.assertEqual(archive.read('tasks/7_example/hebrew.pdf'), b'hebrew')
            self.assertNotIn('tasks/7_example/hebrew.docx', archive.namelist())
            report = json.loads(archive.read('tasks/missing_statements.json'))
            self.assertEqual(report[0]['kind'], 'statement source')
            self.assertEqual(report[0]['digest'], 'source')

    def test_backup_reports_missing_files_and_keeps_other_task_data(self):
        task = self.task()
        healthy = self.task(task_id=8, name='healthy')
        healthy.statements['he'].digest = 'healthy_hebrew'
        dataset = task.active_dataset
        DBTestcase(dataset=dataset, codename='broken_input', input='missing_input', output='output')
        DBTestcase(dataset=dataset, codename='broken_output', input='input', output='missing_output')
        DBTestcase(dataset=dataset, codename='good', input='input', output='output')
        Attachment(task=task, filename='lost.txt', digest='missing_attachment')
        Attachment(task=task, filename='good.txt', digest='good_attachment')
        Manager(dataset=dataset, filename='checker.cpp', digest='missing_manager')
        Generator(dataset=dataset, filename='gen.cpp', digest='missing_generator')
        SubtaskValidator(dataset=dataset, subtask_index=0,
                         filename='val.cpp', digest='missing_validator')
        from datetime import datetime
        submission = Submission(task=task, opaque_id=1, timestamp=datetime(2026, 1, 1),
                                language='Python 3 / CPython')
        File(submission=submission, filename='solution.py', digest='missing_solution')
        ModelSolutionMeta(submission=submission, dataset=dataset, name='model')
        self.missing_file({
            'hebrew', 'missing_input', 'missing_output', 'missing_attachment',
            'missing_manager', 'missing_generator', 'missing_validator', 'missing_solution',
        })
        with self.archive('backup') as archive:
            report = json.loads(archive.read('tasks/missing_files.json'))
            self.assertEqual({entry['kind'] for entry in report}, {
                'statement PDF', 'testcase input', 'testcase output', 'attachment',
                'manager', 'generator', 'subtask validator', 'model solution',
            })
            self.assertEqual(len(report), 8)
            self.assertTrue(all(entry['task_id'] == 7 for entry in report))
            self.assertIn(b'incomplete', archive.read('tasks/INCOMPLETE_BACKUP.txt'))
            base = 'tasks/7_example/'
            self.assertEqual(archive.read(base + 'attachments/good.txt'), b'good_attachment')
            self.assertNotIn(base + 'attachments/lost.txt', archive.namelist())
            self.assertIn(base + 'task.yaml', archive.namelist())
            self.assertIn('tasks/8_healthy/task.yaml', archive.namelist())
            with zipfile.ZipFile(io.BytesIO(archive.read(base + 'tests.zip'))) as tests:
                self.assertEqual(set(tests.namelist()), {
                    'output_broken_input.txt', 'input_broken_output.txt',
                    'input_good.txt', 'output_good.txt',
                })
                self.assertEqual(tests.read('input_good.txt'), b'input')
                self.assertEqual(tests.read('output_broken_input.txt'), b'output')
            missing_input = next(entry for entry in report if entry['kind'] == 'testcase input')
            self.assertEqual(missing_input['testcase'], 'broken_input')
            self.assertEqual(missing_input['digest'], 'missing_input')
        self.assertIn('missing_files.json', self.service.add_notification.call_args.args[2])

    def test_task_packages_remain_strict_by_default(self):
        import tempfile
        from cms.server.admin.handlers.export_handlers import _export_task_to_yaml_format
        task = self.task()
        self.missing_file('hebrew')
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, "task_id.*7.*statement PDF"):
                _export_task_to_yaml_format(task, task.active_dataset,
                                           self.service.file_cacher, directory)

    def test_unknown_export_type(self):
        self.assertEqual(self.fetch('/tasks/export/unknown').code, 404)
